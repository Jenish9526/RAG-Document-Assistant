"""Unit and integration tests for document ingestion, parsing, chunking, and registry."""

import os
import tempfile
import unittest
from io import BytesIO

from src.ingestion.parser import (
    extract_text_from_txt,
    extract_text_from_docx,
    extract_text_from_pdf,
)
from src.ingestion.chunker import split_into_chunks
from src.ingestion.manager import validate_file, process_document, add_document
from src.retrieval.vector_store import VectorStore
from src.utils.helpers import compute_file_hash, clean_text, format_file_size, truncate_text


class TestIngestion(unittest.TestCase):
    def test_utils_helpers(self):
        h1 = compute_file_hash(b"hello world")
        h2 = compute_file_hash(b"hello world")
        h3 = compute_file_hash(b"different")
        self.assertEqual(h1, h2)
        self.assertNotEqual(h1, h3)

        cleaned = clean_text("  hello   world  \n\n\n\ntest  ")
        self.assertIn("hello world", cleaned)
        self.assertNotIn("   ", cleaned)

        self.assertEqual(format_file_size(500), "500.0 B")
        self.assertEqual(format_file_size(2048), "2.0 KB")
        self.assertEqual(format_file_size(5 * 1024 * 1024), "5.0 MB")

        text = "The quick brown fox jumps over the lazy dog"
        truncated = truncate_text(text, max_chars=15)
        self.assertTrue(truncated.endswith("..."))
        self.assertLessEqual(len(truncated), 18)

    def test_extract_text_from_txt(self):
        raw = b"Line 1\nLine 2\nLine 3"
        pages = extract_text_from_txt(raw)
        self.assertEqual(len(pages), 1)
        self.assertEqual(pages[0]["page"], 1)
        self.assertIn("Line 2", pages[0]["text"])

        empty_pages = extract_text_from_txt(b"   \n  ")
        self.assertEqual(len(empty_pages), 0)

    def test_extract_text_from_docx(self):
        import docx
        doc = docx.Document()
        doc.add_paragraph("First paragraph.")
        doc.add_paragraph("Second paragraph.")
        bio = BytesIO()
        doc.save(bio)
        raw = bio.getvalue()

        pages = extract_text_from_docx(raw)
        self.assertEqual(len(pages), 1)
        self.assertIn("First paragraph.", pages[0]["text"])
        self.assertIn("Second paragraph.", pages[0]["text"])

    def test_extract_text_from_pdf(self):
        import pypdf
        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=72, height=72)
        bio = BytesIO()
        writer.write(bio)
        raw = bio.getvalue()

        pages = extract_text_from_pdf(raw)
        self.assertEqual(pages, [])

    def test_split_into_chunks(self):
        words = ["word" + str(i) for i in range(100)]
        text = " ".join(words)
        pages = [{"page": 1, "text": text}]

        chunks = split_into_chunks(pages, document_name="test.txt", chunk_size=30, chunk_overlap=10)
        self.assertGreater(len(chunks), 1)
        for chunk in chunks:
            self.assertEqual(chunk["document"], "test.txt")
            self.assertEqual(chunk["page"], 1)
            self.assertIn("chunk_id", chunk)
            self.assertIn("text", chunk)

    def test_validate_file(self):
        self.assertEqual(validate_file("valid.pdf", b"content"), "")
        self.assertEqual(validate_file("valid.txt", b"content"), "")
        self.assertEqual(validate_file("valid.docx", b"content"), "")

        self.assertIn("Unsupported file type", validate_file("invalid.exe", b"content"))
        self.assertIn("empty", validate_file("empty.txt", b""))

        big_bytes = b"x" * (26 * 1024 * 1024)
        self.assertIn("too large", validate_file("huge.pdf", big_bytes))

    def test_process_document(self):
        res = process_document(b"A quick sample text for processing.", "sample.txt")
        self.assertTrue(res["success"])
        self.assertEqual(res["pages"], 1)
        self.assertGreater(len(res["chunks"]), 0)

        unsupported = process_document(b"content", "file.unknown")
        self.assertFalse(unsupported["success"])

    def test_loader(self):
        from src.ingestion.loader import load_file_from_disk
        with tempfile.NamedTemporaryFile(delete=False, suffix=".txt") as tmp:
            tmp.write(b"sample file on disk")
            tmp_path = tmp.name
        try:
            name, content = load_file_from_disk(tmp_path)
            self.assertEqual(content, b"sample file on disk")
            self.assertTrue(name.endswith(".txt"))
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_manager_lifecycle(self):
        from src.ingestion.manager import clear_all_documents, list_document_names
        with tempfile.TemporaryDirectory() as tmpdir:
            idx_path = os.path.join(tmpdir, "idx.faiss")
            meta_path = os.path.join(tmpdir, "meta.pkl")
            store = VectorStore(index_path=idx_path, metadata_path=meta_path)

            res = add_document("doc.txt", b"Some sample document content.", store=store)
            self.assertTrue(res["success"])
            self.assertEqual(store.total_chunks, 1)

            dup_res = add_document("doc.txt", b"Some sample document content.", store=store)
            self.assertFalse(dup_res["success"])
            self.assertTrue(dup_res["duplicate"])

            clear_all_documents(store=store)
            self.assertEqual(store.total_chunks, 0)


if __name__ == "__main__":
    unittest.main()

