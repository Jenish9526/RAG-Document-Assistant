"""
test_rag_pipeline.py
====================
Comprehensive automated test suite for RAG Document Assistant.
Tests every component:
- utils
- document_processor
- embeddings
- vector_store
- document_manager
- chat_manager
- rag_engine
- llm_service
"""

import os
import sys
import tempfile
import unittest
from io import BytesIO
from unittest.mock import patch, MagicMock

import numpy as np

# Ensure project root is on sys.path for test runners
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

# Component imports
import config
import utils
import document_processor
import embeddings
from vector_store import VectorStore
import document_manager
import chat_manager
import rag_engine
import llm_service


class TestUtils(unittest.TestCase):
    """Test suite for utils.py helper functions."""

    def test_compute_file_hash(self):
        content1 = b"Hello RAG System"
        content2 = b"Hello RAG System"
        content3 = b"Different Content"
        hash1 = utils.compute_file_hash(content1)
        hash2 = utils.compute_file_hash(content2)
        hash3 = utils.compute_file_hash(content3)

        self.assertEqual(len(hash1), 64)
        self.assertEqual(hash1, hash2)
        self.assertNotEqual(hash1, hash3)

    def test_clean_text(self):
        raw = "   This   is   a   test.   \n\n\n\nLine   2 with    spaces.   \n"
        cleaned = utils.clean_text(raw)
        self.assertNotIn("   ", cleaned)
        self.assertNotIn("\n\n\n", cleaned)
        self.assertTrue(cleaned.startswith("This is a test."))
        self.assertTrue(cleaned.endswith("Line 2 with spaces."))
        self.assertEqual(utils.clean_text(""), "")

    def test_format_file_size(self):
        self.assertEqual(utils.format_file_size(500), "500.0 B")
        self.assertEqual(utils.format_file_size(2048), "2.0 KB")
        self.assertEqual(utils.format_file_size(5 * 1024 * 1024), "5.0 MB")
        self.assertEqual(utils.format_file_size(2 * 1024 * 1024 * 1024), "2.0 GB")

    def test_truncate_text(self):
        text = "The quick brown fox jumps over the lazy dog"
        self.assertEqual(utils.truncate_text(text, max_chars=100), text)
        truncated = utils.truncate_text(text, max_chars=20)
        self.assertTrue(truncated.endswith("..."))
        self.assertLessEqual(len(truncated), 23)


class TestDocumentProcessor(unittest.TestCase):
    """Test suite for document extraction and chunking."""

    def test_extract_text_from_txt(self):
        sample_text = "Natural Language Processing with RAG architectures."
        raw_bytes = sample_text.encode("utf-8")
        pages = document_processor.extract_text_from_txt(raw_bytes)
        self.assertEqual(len(pages), 1)
        self.assertEqual(pages[0]["page"], 1)
        self.assertEqual(pages[0]["text"], sample_text)

        # Empty TXT
        empty_pages = document_processor.extract_text_from_txt(b"   \n  ")
        self.assertEqual(empty_pages, [])

    def test_extract_text_from_docx(self):
        import docx
        doc = docx.Document()
        doc.add_paragraph("Paragraph 1 of Word Document.")
        doc.add_paragraph("Paragraph 2 of Word Document.")
        buf = BytesIO()
        doc.save(buf)
        raw_bytes = buf.getvalue()

        pages = document_processor.extract_text_from_docx(raw_bytes)
        self.assertEqual(len(pages), 1)
        self.assertIn("Paragraph 1", pages[0]["text"])
        self.assertIn("Paragraph 2", pages[0]["text"])

    def test_extract_text_from_pdf(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        buf = BytesIO()
        writer.write(buf)
        raw_bytes = buf.getvalue()

        # Blank page should be skipped since it has no extractable text
        pages = document_processor.extract_text_from_pdf(raw_bytes)
        self.assertIsInstance(pages, list)

    def test_split_into_chunks(self):
        # Create 100 words text
        words = [f"word{i}" for i in range(100)]
        text = " ".join(words)
        pages = [{"page": 1, "text": text}]

        chunks = document_processor.split_into_chunks(
            pages, document_name="test.txt", chunk_size=30, chunk_overlap=10
        )
        self.assertGreater(len(chunks), 1)
        for chunk in chunks:
            self.assertEqual(chunk["document"], "test.txt")
            self.assertEqual(chunk["page"], 1)
            self.assertIn("chunk_id", chunk)
            self.assertIn("text", chunk)
            self.assertGreaterEqual(len(chunk["text"]), 10)

    def test_process_document(self):
        # Unsupported extension
        unsupported = document_processor.process_document(b"xyz", "file.png")
        self.assertFalse(unsupported["success"])
        self.assertIn("Unsupported file type", unsupported["error"])

        # Empty content
        empty_doc = document_processor.process_document(b"   ", "empty.txt")
        self.assertFalse(empty_doc["success"])

        # Valid content
        valid_text = ("Artificial Intelligence and Machine Learning " * 20).encode("utf-8")
        res = document_processor.process_document(valid_text, "ai_notes.txt")
        self.assertTrue(res["success"])
        self.assertGreater(res["pages"], 0)
        self.assertGreater(res["characters"], 0)
        self.assertGreater(len(res["chunks"]), 0)


class TestEmbeddings(unittest.TestCase):
    """Test suite for embedding model and vector generation."""

    def test_embeddings_generation(self):
        chunks = [
            {"text": "Retrieval Augmented Generation with FAISS."},
            {"text": "Deep Learning and Transformer Neural Networks."},
        ]
        vectors = embeddings.generate_embeddings(chunks)
        self.assertEqual(vectors.shape, (2, config.EMBEDDING_DIMENSION))
        self.assertEqual(vectors.dtype, np.float32)

        # Check L2 normalization: Euclidean norm of normalized vector should be ~1.0
        norms = np.linalg.norm(vectors, axis=1)
        for norm in norms:
            self.assertAlmostEqual(float(norm), 1.0, places=4)

    def test_query_embedding_generation(self):
        query = "What is RAG?"
        vector = embeddings.generate_query_embedding(query)
        self.assertEqual(vector.shape, (1, config.EMBEDDING_DIMENSION))
        self.assertEqual(vector.dtype, np.float32)
        norm = np.linalg.norm(vector[0])
        self.assertAlmostEqual(float(norm), 1.0, places=4)


class TestVectorStore(unittest.TestCase):
    """Test suite for FAISS vector store operations."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.idx_path = os.path.join(self.temp_dir.name, "test.faiss")
        self.meta_path = os.path.join(self.temp_dir.name, "test_meta.pkl")
        self.store = VectorStore(index_path=self.idx_path, metadata_path=self.meta_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_add_and_search(self):
        chunks = [
            {"document": "doc1.txt", "page": 1, "chunk_id": 0, "text": "Python programming language"},
            {"document": "doc2.txt", "page": 1, "chunk_id": 1, "text": "Cooking Italian pasta recipes"},
        ]
        vectors = embeddings.generate_embeddings(chunks)
        self.store.add_documents(chunks, vectors)

        self.assertEqual(self.store.total_chunks, 2)
        self.assertEqual(len(self.store.chunks_for_document("doc1.txt")), 1)

        # Query relevant to Python
        q_vec = embeddings.generate_query_embedding("writing python software code")
        results = self.store.search(q_vec, top_k=2)

        self.assertEqual(len(results), 2)
        top_match_chunk, top_score = results[0]
        self.assertEqual(top_match_chunk["document"], "doc1.txt")
        self.assertGreater(top_score, results[1][1])

    def test_persistence_save_load(self):
        chunks = [{"document": "d.txt", "page": 1, "chunk_id": 0, "text": "persisted content"}]
        vectors = embeddings.generate_embeddings(chunks)
        self.store.add_documents(chunks, vectors)
        self.store.save_index()

        # Load into new store instance with same paths
        new_store = VectorStore(index_path=self.idx_path, metadata_path=self.meta_path)
        loaded = new_store.load_index()
        self.assertTrue(loaded)
        self.assertEqual(new_store.total_chunks, 1)
        self.assertEqual(new_store.metadata[0]["text"], "persisted content")

    def test_clear_index(self):
        chunks = [{"document": "d.txt", "page": 1, "chunk_id": 0, "text": "sample"}]
        vectors = embeddings.generate_embeddings(chunks)
        self.store.add_documents(chunks, vectors)
        self.assertEqual(self.store.total_chunks, 1)

        self.store.clear_index()
        self.assertEqual(self.store.total_chunks, 0)
        self.assertEqual(len(self.store.metadata), 0)


class TestDocumentManager(unittest.TestCase):
    """Test suite for document validation and registry management."""

    def test_validate_file(self):
        # Valid files
        self.assertEqual(document_manager.validate_file("notes.pdf", b"pdfcontent"), "")
        self.assertEqual(document_manager.validate_file("notes.txt", b"txtcontent"), "")
        self.assertEqual(document_manager.validate_file("notes.docx", b"docxcontent"), "")

        # Invalid extension
        err_ext = document_manager.validate_file("notes.exe", b"content")
        self.assertIn("Unsupported file type", err_ext)

        # Empty file
        err_empty = document_manager.validate_file("notes.txt", b"")
        self.assertIn("empty", err_empty.lower())

        # Oversized file
        huge_bytes = b"0" * ((config.MAX_FILE_SIZE_MB + 1) * 1024 * 1024)
        err_size = document_manager.validate_file("huge.txt", huge_bytes)
        self.assertIn("too large", err_size.lower())


class TestChatManager(unittest.TestCase):
    """Test suite for in-memory chat history."""

    def test_chat_lifecycle(self):
        chat_manager.init_chat_history()
        chat_manager.clear_history()
        self.assertEqual(chat_manager.get_history(), [])

        chat_manager.add_message("user", "Hello Assistant")
        chat_manager.add_message("assistant", "Hello! How can I help?", sources=[{"document": "test.txt", "page": 1, "score": 0.85}])

        history = chat_manager.get_history()
        self.assertEqual(len(history), 2)
        self.assertEqual(history[0]["role"], "user")
        self.assertEqual(history[0]["content"], "Hello Assistant")
        self.assertEqual(history[1]["role"], "assistant")
        self.assertEqual(len(history[1]["sources"]), 1)

        chat_manager.clear_history()
        self.assertEqual(chat_manager.get_history(), [])


class TestRagEngine(unittest.TestCase):
    """Test suite for RAG context building, prompts, and retrieval."""

    def test_build_context(self):
        retrieved = [
            ({"document": "docA.pdf", "page": 2, "text": "Content of chunk 1."}, 0.82),
            ({"document": "docB.pdf", "page": 5, "text": "Content of chunk 2."}, 0.74),
        ]
        context = rag_engine.build_context(retrieved)
        self.assertIn("[Source 1: docA.pdf, Page 2]", context)
        self.assertIn("Content of chunk 1.", context)
        self.assertIn("[Source 2: docB.pdf, Page 5]", context)

    def test_build_prompt_variants(self):
        query = "Explain binary search."
        context = "[Source 1: algo.txt, Page 1]\nBinary search runs in O(log n)."

        # Simple prompt
        prompt_simple = rag_engine.build_prompt(query, context, answer_style="Simple", exam_mode=False)
        self.assertIn("easy language", prompt_simple)
        self.assertIn(query, prompt_simple)
        self.assertIn(context, prompt_simple)

        # Detailed prompt
        prompt_detailed = rag_engine.build_prompt(query, context, answer_style="Detailed", exam_mode=False)
        self.assertIn("more thorough explanation", prompt_detailed)

        # Exam mode prompt
        prompt_exam = rag_engine.build_prompt(query, context, answer_style="Detailed", exam_mode=True)
        self.assertIn("Definition, Explanation, Steps, Example", prompt_exam)

    def test_answer_question_no_docs(self):
        empty_store = VectorStore()
        res = rag_engine.answer_question("Any question?", empty_store)
        self.assertFalse(res["found_context"])
        self.assertIn("upload a document", res["answer"].lower())

    def test_answer_question_empty_query(self):
        store = VectorStore()
        res = rag_engine.answer_question("   ", store)
        self.assertFalse(res["found_context"])
        self.assertIn("enter a question", res["answer"].lower())

    @patch("llm_service.generate_response")
    def test_answer_question_end_to_end(self, mock_llm):
        mock_llm.return_value = "Binary search is an efficient algorithm with O(log n) complexity."

        store = VectorStore()
        chunks = [
            {
                "document": "dsa.txt",
                "page": 1,
                "chunk_id": 0,
                "text": "Binary search is an algorithm for finding an element in a sorted list.",
            }
        ]
        vectors = embeddings.generate_embeddings(chunks)
        store.add_documents(chunks, vectors)

        res = rag_engine.answer_question("What is binary search?", store, top_k=1)
        self.assertTrue(res["found_context"])
        self.assertIn("Binary search", res["answer"])
        self.assertEqual(len(res["sources"]), 1)
        self.assertEqual(res["sources"][0]["document"], "dsa.txt")

    @patch("llm_service.generate_response")
    def test_summarize_document(self, mock_llm):
        mock_llm.return_value = "Summary: Main Topic: DSA, Key Points: Algorithms."

        store = VectorStore()
        chunks = [
            {"document": "dsa.txt", "page": 1, "chunk_id": 0, "text": "Data structures and algorithms overview."}
        ]
        store.metadata.extend(chunks)

        summary = rag_engine.summarize_document("dsa.txt", store)
        self.assertIn("Summary", summary)

    def test_generate_suggested_questions_fallback(self):
        store = VectorStore()
        # With no LLM configured / no chunks, returns fallback questions
        with patch("llm_service.is_configured", return_value=False):
            questions = rag_engine.generate_suggested_questions("doc.txt", store)
            self.assertEqual(len(questions), 5)
            self.assertIn("What is the main topic of this document?", questions)


class TestLLMService(unittest.TestCase):
    """Test suite for LLM service configuration and error handling."""

    def test_unconfigured_error(self):
        with patch.object(llm_service, "LLM_API_KEY", ""):
            self.assertFalse(llm_service.is_configured())
            with self.assertRaises(llm_service.LLMNotConfiguredError):
                llm_service.generate_response("Test prompt")


if __name__ == "__main__":
    unittest.main()
